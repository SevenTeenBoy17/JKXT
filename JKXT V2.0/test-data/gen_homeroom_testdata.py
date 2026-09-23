# -*- coding: utf-8 -*-
"""第4次测试数据 — 验证"班主任不监本班"(1.1)。
场景：高一 6 个班，每班 1 名班主任；4 科考试，每班在自己考场考，每场绑定该班。
列(对齐 1.1 新增):
  教师: 工号 姓名 性别 年级 任教学科 手机号 邮箱 班主任 所带班级
  班级: 班级编号 班级名称 年级 学段 学生数 开设科目 班主任工号
  考场: 考场编号 考场名称 楼栋 楼层 容量 考试容量 考场类型 设施标签
  考试: 考试科目 年级 考试日期 开始时间 结束时间 考场编号 人数 监考人数 班级
"""
import os, random
import openpyxl

random.seed(20260619)
SUR = list("王李张刘陈杨赵黄周吴徐孙马朱胡郭何高林罗郑梁谢宋唐许韩冯邓曹彭曾肖田董袁潘于蒋蔡余杜叶程苏魏吕丁任沈姚卢")
GIV = list("伟芳娜敏静秀英华慧强磊洋勇艳杰娟涛明超霞平刚桂兰建文军婷雪琳云宇辰睿欣怡浩然轩涵远宁航曦")
def nm(): return random.choice(SUR) + random.choice(GIV) + (random.choice(GIV) if random.random() < .4 else "")

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "04-班主任避本班")
os.makedirs(OUT, exist_ok=True)

NCLASS = 6
SUBJECTS = ["语文", "数学", "英语", "物理"]
SLOTS = [("2026-06-15", "08:00", "10:00"), ("2026-06-15", "10:30", "12:00"),
         ("2026-06-16", "08:00", "10:00"), ("2026-06-16", "10:30", "12:00")]

# --- 教师: 6 班主任(T班01..) + 18 普通(T26001..)，全高一 ---
tw = openpyxl.Workbook(); ts = tw.active; ts.title = "教师表"
ts.append(["工号", "姓名", "性别", "年级", "任教学科", "手机号", "邮箱", "班主任", "所带班级"])
homeroom_no = {}   # class_code -> teacher_no
for c in range(1, NCLASS + 1):
    no = "TBZ%02d" % c; cls = "G1-%02d" % c; homeroom_no[cls] = no
    subj = SUBJECTS[(c - 1) % len(SUBJECTS)]
    ts.append([no, nm(), random.choice(["男", "女"]), "高一", subj,
               "138%08d" % random.randint(0, 99999999), "%s@tjsyzx.cn" % no.lower(),
               "是", cls])   # 教师表也填班主任+所带班级(双路径之一)
for i in range(1, 19):
    no = "T26%03d" % i
    ts.append([no, nm(), random.choice(["男", "女"]), "高一", SUBJECTS[i % len(SUBJECTS)],
               "139%08d" % random.randint(0, 99999999), "%s@tjsyzx.cn" % no.lower(), "否", ""])
tw.save(os.path.join(OUT, "01-教师安排-第4次-班主任避本班.xlsx"))

# --- 班级: 6 个，班主任工号(另一条路径) ---
cw = openpyxl.Workbook(); cs = cw.active; cs.title = "班级表"
cs.append(["班级编号", "班级名称", "年级", "学段", "学生数", "开设科目", "班主任工号"])
for c in range(1, NCLASS + 1):
    cls = "G1-%02d" % c
    cs.append([cls, "高一(%d)班" % c, "高一", "高中", 45, ",".join(SUBJECTS), homeroom_no[cls]])
cw.save(os.path.join(OUT, "02-班级信息-第4次-班主任避本班.xlsx"))

# --- 考场: 6 间 ---
rw = openpyxl.Workbook(); rs = rw.active; rs.title = "考场表"
rs.append(["考场编号", "考场名称", "楼栋", "楼层", "容量", "考试容量", "考场类型", "设施标签"])
for c in range(1, NCLASS + 1):
    rs.append(["R10%d" % c, "明德楼 10%d 教室" % c, "明德楼", "1F", 50, 48, "standard", "广播,监控,空调"])
rw.save(os.path.join(OUT, "03-考场信息-第4次-班主任避本班.xlsx"))

# --- 考试: 每班×每科 = 24 场；班 c 的 4 科分到 4 个时段，固定在 R10c，绑定班级 G1-0c ---
ew = openpyxl.Workbook(); es = ew.active; es.title = "考试表"
es.append(["考试科目", "年级", "考试日期", "开始时间", "结束时间", "考场编号", "人数", "监考人数", "班级"])
for c in range(1, NCLASS + 1):
    cls = "G1-%02d" % c; room = "R10%d" % c
    for si, subj in enumerate(SUBJECTS):
        d, s, e = SLOTS[si]
        es.append([subj, "高一", d, s, e, room, 45, 2, cls])
ew.save(os.path.join(OUT, "04-考试安排-第4次-班主任避本班.xlsx"))

print("OK 第4次: classes=%d teachers=%d(6班主任+18) rooms=%d exams=%d" % (NCLASS, 6 + 18, NCLASS, NCLASS * len(SUBJECTS)))
